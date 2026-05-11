#!/usr/bin/env python3
"""Per-stock PPT builder — 종목당 10-12 slide deep-dive deck.

build_overview_pptx_v2.py와 동일한 design system 재사용:
* 16:9 master template (chapter ribbon · title · subtitle · body 일관 위치)
* Forest Canopy theme (deep forest + gold + cream)
* McKinsey pyramid (top message first)
* 매 slide visual element (chart/badge/icon)
* AI-slop 회피

Slides:
  1. Cover (deep forest)
  2. Investment Snapshot (Live market data 카드)
  3. Thesis Alignment (블로거 thesis ↔ 종목 fit)
  4. 13명 페르소나 verdict
  5. Style split (가치/성장/매크로 등)
  6. Universal concerns + opportunities
  7. Bull / Base / Bear scenarios
  8. Catalysts timeline
  9. Risks matrix
  10. Industry context + competitors
  11. Decision sheet (Risk Manager + Portfolio Manager + 결론)

Usage:
  python build_per_stock_pptx.py \\
    --pipeline-dir .analysis-log/bloggers/{blogger}/{slug} \\
    --ticker LYC.AX \\
    --output 1_LYC.AX_라이너스.pptx
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

SCRIPT_DIR = Path(__file__).resolve().parent
PLUGIN_ROOT = SCRIPT_DIR.parent.parent.parent
sys.path.insert(0, str(PLUGIN_ROOT / "skills" / "_common"))
from ticker_resolver import resolve_ticker  # type: ignore

# Reuse design system from overview_pptx_v2
sys.path.insert(0, str(SCRIPT_DIR))
from build_overview_pptx_v2 import (  # type: ignore
    DEEP_FOREST, MOSS, SAGE, GOLD, CREAM, INK, GRAY, LIGHT_GRAY,
    BULL, NEU, BEAR, WHITE, HEADER_FONT, BODY_FONT,
    SLIDE_W, SLIDE_H, CHAPTER_TOP, CHAPTER_LEFT, CHAPTER_W, CHAPTER_H,
    TITLE_TOP, TITLE_LEFT, TITLE_W, TITLE_H,
    SUB_TOP, SUB_LEFT, SUB_W, SUB_H,
    BODY_TOP, BODY_LEFT, BODY_W, BODY_H,
    FOOTER_TOP, FOOTER_LEFT, FOOTER_W, FOOTER_H,
    add_text, add_rect, add_circle, add_chapter_ribbon, add_title,
    add_footer, base_slide, _read_json,
)


# ── Slide builders ────────────────────────────────────────────────────────
def cover_slide(prs, ticker, company, market, score, action):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = DEEP_FOREST
    bg.line.fill.background()

    # accent
    add_rect(slide, Inches(11.5), Inches(0.6), Inches(1.3), Inches(0.08), GOLD)
    add_circle(slide, Inches(11.5) - Inches(0.14), Inches(0.66) - Inches(0.04),
               Inches(0.16), GOLD)

    # category
    add_text(slide, Inches(0.7), Inches(0.7), Inches(8), Inches(0.4),
             "INVESTMENT THESIS · DEEP DIVE", size=10, color=GOLD,
             bold=True, font=BODY_FONT)

    # ticker (large mono)
    add_text(slide, Inches(0.7), Inches(1.7), Inches(11.9), Inches(0.6),
             ticker, size=22, color=SAGE, bold=True, font=HEADER_FONT)

    # company (huge)
    add_text(slide, Inches(0.7), Inches(2.3), Inches(11.9), Inches(1.4),
             company, size=44, bold=True, color=WHITE, font=HEADER_FONT)

    # tagline
    info = resolve_ticker(ticker)
    add_text(slide, Inches(0.7), Inches(3.9), Inches(11.9), Inches(0.7),
             info.get("sector", ""), size=14, color=SAGE, italic=True,
             font=BODY_FONT)

    # divider
    add_rect(slide, Inches(0.7), Inches(4.9), Inches(2.0), Inches(0.04), GOLD)

    # 4 KPI strip
    kpis = [
        ("Score (페르소나)", f"{score:.2f}/5"),
        ("Action", action),
        ("현재가", f"{market.get('current_price', 0):,.2f} {market.get('currency', '')}"),
        ("1Y Return", f"+{market.get('return_1y_pct', 0):.0f}%"),
    ]
    kpi_top = Inches(5.4)
    kpi_w = Inches(2.95)
    for i, (lbl, val) in enumerate(kpis):
        kx = Inches(0.7) + (kpi_w + Inches(0.05)) * i
        add_text(slide, kx, kpi_top, kpi_w, Inches(0.30),
                 lbl, size=10, color=SAGE, font=BODY_FONT)
        add_text(slide, kx, kpi_top + Inches(0.30), kpi_w, Inches(0.55),
                 val, size=22, bold=True, color=GOLD, font=HEADER_FONT)

    # Footer
    add_text(slide, Inches(0.7), Inches(7.05), Inches(11.9), Inches(0.30),
             f"분석일 2026-05-07 · 13명 페르소나 + Deep Research · 블로거 dhgusdnd44 (happy)",
             size=9, color=SAGE, font=BODY_FONT)
    return slide


def snapshot_slide(prs, ticker, company, market, score, page, total):
    """Slide 2 — Live market snapshot 6 KPI cards."""
    slide = base_slide(prs, "01 · Investment Snapshot",
                       f"{company} ({ticker}) — yfinance live 2026-05-07",
                       f"시총 ${market.get('market_cap_usd_bn', 0):.1f}B · fwd PE {market.get('forward_pe', 0):.1f} · vol {market.get('volatility_annualized_pct', 0):.0f}% · β {market.get('beta', 0):.2f}",
                       page, total)
    # 2x3 KPI grid
    kpis = [
        ("현재가", f"{market.get('current_price', 0):,.2f}", market.get('currency', ''), DEEP_FOREST),
        ("1Y Return", f"+{market.get('return_1y_pct', 0):.0f}%", "1년 수익률", BULL if market.get('return_1y_pct', 0) > 0 else BEAR),
        ("3M Return", f"+{market.get('return_3m_pct', 0):.0f}%", "3개월 수익률",
         BULL if market.get('return_3m_pct', 0) > 0 else BEAR),
        ("Vol annualized", f"{market.get('volatility_annualized_pct', 0):.0f}%", "연환산 변동성",
         BEAR if market.get('volatility_annualized_pct', 0) > 80 else (GOLD if market.get('volatility_annualized_pct', 0) > 50 else BULL)),
        ("Forward PE", f"{market.get('forward_pe', 0):.1f}x" if market.get('forward_pe', 0) > 0 else "적자", "선행 PER",
         BEAR if market.get('forward_pe', 0) > 50 else (GOLD if market.get('forward_pe', 0) > 25 else BULL)),
        ("Beta", f"{market.get('beta', 0):.2f}", "시장 민감도", DEEP_FOREST),
    ]
    card_w = Inches(4.0)
    card_h = Inches(2.05)
    gap = Inches(0.15)
    for i, (lbl, val, sub, c) in enumerate(kpis):
        row, col = i // 3, i % 3
        left = BODY_LEFT + col * (card_w + gap)
        top = BODY_TOP + Inches(0.10) + row * (card_h + gap)
        add_rect(slide, left, top, card_w, card_h, WHITE, line=LIGHT_GRAY)
        # accent left bar
        add_rect(slide, left, top, Inches(0.10), card_h, c)
        add_text(slide, left + Inches(0.30), top + Inches(0.15),
                 card_w - Inches(0.40), Inches(0.30),
                 lbl, size=11, color=GRAY, font=BODY_FONT)
        add_text(slide, left + Inches(0.30), top + Inches(0.50),
                 card_w - Inches(0.40), Inches(0.95),
                 val, size=42, bold=True, color=c, font=HEADER_FONT)
        add_text(slide, left + Inches(0.30), top + Inches(1.55),
                 card_w - Inches(0.40), Inches(0.30),
                 sub, size=10, color=GRAY, italic=True, font=BODY_FONT)


def thesis_alignment_slide(prs, deep_research, page, total):
    """Slide 3 — 블로거 thesis ↔ 종목 fit."""
    company = deep_research.get('company', '')
    overview = deep_research.get('industry', {}).get('overview', '')[:300]
    slide = base_slide(prs, "02 · Thesis Alignment",
                       "블로거 thesis와 종목 fit 분석",
                       overview[:120],
                       page, total)
    # Industry overview block
    ov_top = BODY_TOP + Inches(0.10)
    add_rect(slide, BODY_LEFT, ov_top, BODY_W, Inches(1.50), WHITE,
             line=LIGHT_GRAY)
    add_rect(slide, BODY_LEFT, ov_top, Inches(0.10), Inches(1.50), GOLD)
    add_text(slide, BODY_LEFT + Inches(0.30), ov_top + Inches(0.15),
             BODY_W - Inches(0.40), Inches(0.30),
             "🔍 산업 overview", size=11, bold=True, color=DEEP_FOREST,
             font=HEADER_FONT)
    add_text(slide, BODY_LEFT + Inches(0.30), ov_top + Inches(0.50),
             BODY_W - Inches(0.40), Inches(0.95),
             overview, size=11, color=INK, font=BODY_FONT)

    # Thesis fit boxes (3 main thesis 매핑)
    fit_top = ov_top + Inches(1.70)
    fit_h = Inches(2.55)
    col_w = (BODY_W - Inches(0.40)) / 3
    fits = [
        ("T1 정제 병목론", "13/13 정확 fit", "비중국 정제 vertical에 직접 노출", BULL),
        ("T6 중희토류 비대칭", "10/13 부분 fit", "Dy/Tb capacity 진입 timeline", GOLD),
        ("T8 광산-정제 mispricing", "13/13 alpha source", "Damodaran·Ackman의 핵심 thesis", BULL),
    ]
    for i, (tid, score_label, desc, c) in enumerate(fits):
        left = BODY_LEFT + (col_w + Inches(0.20)) * i
        # use int to avoid Inches arithmetic issues
        left_emu = int(BODY_LEFT) + (int(col_w) + int(Inches(0.20))) * i
        add_rect(slide, Emu(left_emu), fit_top, Emu(int(col_w)), fit_h, WHITE,
                 line=LIGHT_GRAY)
        add_rect(slide, Emu(left_emu), fit_top, Emu(int(col_w)), Inches(0.40), c)
        add_text(slide, Emu(left_emu), fit_top, Emu(int(col_w)), Inches(0.40),
                 tid, size=12, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=BODY_FONT)
        add_text(slide, Emu(left_emu + int(Inches(0.20))),
                 fit_top + Inches(0.55), Emu(int(col_w) - int(Inches(0.40))),
                 Inches(0.40), score_label, size=20, bold=True, color=c,
                 font=HEADER_FONT)
        add_text(slide, Emu(left_emu + int(Inches(0.20))),
                 fit_top + Inches(1.05), Emu(int(col_w) - int(Inches(0.40))),
                 Inches(1.40), desc, size=10, color=INK, italic=True,
                 font=BODY_FONT)


def persona_verdict_slide(prs, persona_agg, page, total):
    """Slide 4 — 13명 verdict 분포."""
    score = persona_agg.get('weighted_average_score_1to5', 0)
    consensus = persona_agg.get('consensus_verdict', '')[:60]
    slide = base_slide(prs, "03 · Persona Panel",
                       f"13명 가중 score {score:.2f}/5 — {consensus}",
                       f"평균 신뢰도 {int(persona_agg.get('average_confidence', 0)*100)}% · Buffett·Munger·Lynch·Wood·Burry·Taleb·Graham·Ackman·Pabrai·Fisher·Jhunjhunwala·Druckenmiller·Damodaran",
                       page, total)
    dist = persona_agg.get('verdict_distribution', {})
    # Stacked horizontal bar
    bar_left = BODY_LEFT
    bar_top = BODY_TOP + Inches(0.30)
    bar_total_w = BODY_W
    bar_h = Inches(1.20)
    add_text(slide, bar_left, bar_top - Inches(0.30), bar_total_w, Inches(0.25),
             "Verdict 분포 (총 13명)", size=11, bold=True, color=DEEP_FOREST,
             font=HEADER_FONT)
    cats = [
        ("very_bullish", "Very Bullish", BULL),
        ("lean_bullish", "Lean Bullish", RGBColor(0x66, 0xBB, 0x6A)),
        ("neutral", "Neutral", NEU),
        ("lean_bearish", "Lean Bearish", RGBColor(0xEF, 0x53, 0x50)),
        ("very_bearish", "Very Bearish", BEAR),
    ]
    cur_x = bar_left
    for key, lbl, c in cats:
        cnt = dist.get(key, {}).get('count', 0)
        if cnt == 0:
            continue
        seg_w = Emu(int(int(bar_total_w) * cnt / 13))
        add_rect(slide, cur_x, bar_top, seg_w, bar_h, c)
        add_text(slide, cur_x, bar_top + Inches(0.20), seg_w, Inches(0.35),
                 f"{cnt}명", size=14, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=HEADER_FONT)
        add_text(slide, cur_x, bar_top + Inches(0.65), seg_w, Inches(0.35),
                 lbl, size=9, color=WHITE, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, font=BODY_FONT)
        cur_x += seg_w

    # 13 persona dots row (visual reinforcement)
    persona_results = persona_agg.get('persona_results', persona_agg.get('persona_results_summary', {}))
    dot_top = bar_top + bar_h + Inches(0.50)
    add_text(slide, BODY_LEFT, dot_top - Inches(0.30), BODY_W, Inches(0.25),
             "13명 개별 verdict (왼쪽: Buffett → 오른쪽: Damodaran)",
             size=11, bold=True, color=DEEP_FOREST, font=HEADER_FONT)
    persona_order = ["warren-buffett", "charlie-munger", "ben-graham",
                     "mohnish-pabrai", "michael-burry", "peter-lynch",
                     "cathie-wood", "phil-fisher", "stanley-druckenmiller",
                     "bill-ackman", "nassim-taleb", "rakesh-jhunjhunwala",
                     "aswath-damodaran"]
    persona_short = ["Buf", "Mun", "Gra", "Pab", "Bur", "Lyn", "Wod",
                     "Fis", "Dru", "Ack", "Tal", "Jhu", "Dam"]
    verdict_color = {
        "very_bullish": BULL, "lean_bullish": RGBColor(0x66, 0xBB, 0x6A),
        "neutral": NEU, "lean_bearish": RGBColor(0xEF, 0x53, 0x50),
        "very_bearish": BEAR,
    }
    dot_d = Inches(0.85)
    gap_x = Inches(0.10)
    for i, pid in enumerate(persona_order):
        x = BODY_LEFT + (dot_d + gap_x) * i
        pr = persona_results.get(pid, {})
        v = pr.get('verdict', 'neutral')
        c = verdict_color.get(v, NEU)
        add_circle(slide, x, dot_top, dot_d, c)
        add_text(slide, x, dot_top, dot_d, dot_d, persona_short[i], size=11,
                 bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, font=BODY_FONT)
        # confidence below
        conf = pr.get('confidence', 0)
        add_text(slide, x, dot_top + dot_d + Inches(0.05), dot_d, Inches(0.20),
                 f"{int(conf*100)}%", size=8, color=GRAY, align=PP_ALIGN.CENTER,
                 font=BODY_FONT)


def style_split_slide(prs, persona_agg, page, total):
    """Slide 5 — Style 카테고리별 verdict (가치/성장/매크로 등)."""
    slide = base_slide(prs, "04 · Style Split",
                       "투자 스타일별 verdict 분포",
                       "가치파(Buffett·Munger·Graham·Pabrai) vs 성장파(Lynch·Wood·Fisher) vs 매크로(Druckenmiller) 등",
                       page, total)
    style_split = persona_agg.get('style_split', {})
    # 8 style cards (2x4)
    styles = [
        ("value", "가치파", "Buffett·Munger·Graham·Pabrai"),
        ("value-contrarian", "Contrarian", "Burry"),
        ("growth", "성장파", "Lynch·Wood·Fisher"),
        ("macro", "매크로", "Druckenmiller"),
        ("activist", "Activist", "Ackman"),
        ("risk", "Risk/Tail", "Taleb"),
        ("growth-em", "EM Growth", "Jhunjhunwala"),
        ("valuation", "DCF/Valuation", "Damodaran"),
    ]
    card_w = Inches(2.95)
    card_h = Inches(2.10)
    gap = Inches(0.15)
    for i, (key, lbl, members) in enumerate(styles):
        row, col = i // 4, i % 4
        left_emu = int(BODY_LEFT) + col * (int(card_w) + int(gap))
        top_emu = int(BODY_TOP + Inches(0.10)) + row * (int(card_h) + int(gap))
        info = style_split.get(key, {})
        dom = info.get('dominant', 'no data')[:50]
        # determine card accent color from dominant
        if 'very_bullish' in dom:
            c = BULL
        elif 'lean_bullish' in dom:
            c = RGBColor(0x66, 0xBB, 0x6A)
        elif 'lean_bearish' in dom:
            c = RGBColor(0xEF, 0x53, 0x50)
        elif 'very_bearish' in dom:
            c = BEAR
        else:
            c = NEU
        add_rect(slide, Emu(left_emu), Emu(top_emu), card_w, card_h, WHITE,
                 line=LIGHT_GRAY)
        add_rect(slide, Emu(left_emu), Emu(top_emu), card_w, Inches(0.35), c)
        add_text(slide, Emu(left_emu), Emu(top_emu), card_w, Inches(0.35),
                 lbl, size=12, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=BODY_FONT)
        add_text(slide, Emu(left_emu + int(Inches(0.15))),
                 Emu(top_emu + int(Inches(0.45))),
                 card_w - Inches(0.30), Inches(0.30),
                 members, size=9, color=GRAY, italic=True, font=BODY_FONT)
        add_text(slide, Emu(left_emu + int(Inches(0.15))),
                 Emu(top_emu + int(Inches(0.85))),
                 card_w - Inches(0.30), Inches(1.15),
                 dom, size=11, bold=True, color=c, font=BODY_FONT)


def concerns_opportunities_slide(prs, persona_agg, page, total):
    """Slide 6 — Universal concerns & opportunities."""
    slide = base_slide(prs, "05 · Concerns & Opportunities",
                       "13명 페르소나 빈도 가중 핵심 우려·기회",
                       "Universal concerns (다수 페르소나가 공통 지적) vs Universal opportunities",
                       page, total)
    concerns = persona_agg.get('universal_concerns', [])[:4]
    opps = persona_agg.get('universal_opportunities', [])[:4]

    col_w = (BODY_W - Inches(0.30)) / 2
    # Left: concerns
    left_x = BODY_LEFT
    add_text(slide, left_x, BODY_TOP + Inches(0.10), col_w, Inches(0.30),
             "⚠ 핵심 우려 (Universal Concerns)", size=12, bold=True,
             color=BEAR, font=HEADER_FONT)
    for i, c in enumerate(concerns):
        ttop = BODY_TOP + Inches(0.50) + Inches(0.95) * i
        add_rect(slide, left_x, ttop, col_w, Inches(0.85), WHITE,
                 line=LIGHT_GRAY)
        add_rect(slide, left_x, ttop, Inches(0.10), Inches(0.85), BEAR)
        add_text(slide, left_x + Inches(0.25), ttop + Inches(0.10),
                 col_w - Inches(0.40), Inches(0.30),
                 c.get('theme', '')[:60], size=11, bold=True, color=DEEP_FOREST,
                 font=HEADER_FONT)
        ex = c.get('examples', [''])[0][:130]
        add_text(slide, left_x + Inches(0.25), ttop + Inches(0.40),
                 col_w - Inches(0.40), Inches(0.40),
                 ex, size=9, color=INK, font=BODY_FONT)
        # frequency badge
        freq = c.get('frequency', 0)
        add_circle(slide, left_x + col_w - Inches(0.50), ttop + Inches(0.20),
                   Inches(0.45), BEAR)
        add_text(slide, left_x + col_w - Inches(0.50), ttop + Inches(0.20),
                 Inches(0.45), Inches(0.45), f"{freq}/13", size=10, bold=True,
                 color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=BODY_FONT)

    # Right: opportunities
    right_x = BODY_LEFT + col_w + Inches(0.30)
    add_text(slide, right_x, BODY_TOP + Inches(0.10), col_w, Inches(0.30),
             "✓ 핵심 기회 (Universal Opportunities)", size=12, bold=True,
             color=BULL, font=HEADER_FONT)
    for i, o in enumerate(opps):
        ttop = BODY_TOP + Inches(0.50) + Inches(0.95) * i
        add_rect(slide, right_x, ttop, col_w, Inches(0.85), WHITE,
                 line=LIGHT_GRAY)
        add_rect(slide, right_x, ttop, Inches(0.10), Inches(0.85), BULL)
        add_text(slide, right_x + Inches(0.25), ttop + Inches(0.10),
                 col_w - Inches(0.40), Inches(0.30),
                 o.get('theme', '')[:60], size=11, bold=True, color=DEEP_FOREST,
                 font=HEADER_FONT)
        ex = o.get('examples', [''])[0][:130]
        add_text(slide, right_x + Inches(0.25), ttop + Inches(0.40),
                 col_w - Inches(0.40), Inches(0.40),
                 ex, size=9, color=INK, font=BODY_FONT)
        freq = o.get('frequency', 0)
        add_circle(slide, right_x + col_w - Inches(0.50), ttop + Inches(0.20),
                   Inches(0.45), BULL)
        add_text(slide, right_x + col_w - Inches(0.50), ttop + Inches(0.20),
                 Inches(0.45), Inches(0.45), f"{freq}/13", size=10, bold=True,
                 color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=BODY_FONT)


def scenarios_slide(prs, deep_research, market, page, total):
    """Slide 7 — Bull/Base/Bear."""
    cur_price = market.get('current_price', 0)
    cur_curr = market.get('currency', '')
    sc = deep_research.get('scenarios', {})
    bull, base, bear = sc.get('bull', {}), sc.get('base', {}), sc.get('bear', {})

    # Calculate expected
    expected = (bull.get('prob_pct', 0) * bull.get('target_price', 0) +
                base.get('prob_pct', 0) * base.get('target_price', 0) +
                bear.get('prob_pct', 0) * bear.get('target_price', 0)) / 100.0
    upside = (expected / cur_price - 1) * 100 if cur_price else 0

    slide = base_slide(prs, "06 · Scenarios",
                       f"확률 가중 기대 주가 {expected:,.2f} {cur_curr} ({upside:+.1f}%)",
                       f"현재가 {cur_price:,.2f} {cur_curr} · Bull {bull.get('prob_pct', 0)}% / Base {base.get('prob_pct', 0)}% / Bear {bear.get('prob_pct', 0)}%",
                       page, total)
    # 3 scenario columns
    col_w = (BODY_W - Inches(0.40)) / 3
    scenarios = [
        ("Bull", bull, BULL),
        ("Base", base, NEU),
        ("Bear", bear, BEAR),
    ]
    for i, (name, sd, c) in enumerate(scenarios):
        left_emu = int(BODY_LEFT) + i * (int(col_w) + int(Inches(0.20)))
        top = BODY_TOP + Inches(0.10)
        h = Inches(4.55)
        add_rect(slide, Emu(left_emu), top, Emu(int(col_w)), h, WHITE,
                 line=LIGHT_GRAY)
        # header
        add_rect(slide, Emu(left_emu), top, Emu(int(col_w)), Inches(0.50), c)
        add_text(slide, Emu(left_emu), top, Emu(int(col_w)), Inches(0.50),
                 f"{name}  ·  {sd.get('prob_pct', 0)}%", size=14, bold=True,
                 color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=HEADER_FONT)
        # target price huge
        target = sd.get('target_price', 0)
        scenario_upside = (target / cur_price - 1) * 100 if cur_price else 0
        add_text(slide, Emu(left_emu + int(Inches(0.20))),
                 top + Inches(0.65), Emu(int(col_w) - int(Inches(0.40))),
                 Inches(0.65),
                 f"{target:,.2f}", size=42, bold=True, color=DEEP_FOREST,
                 font=HEADER_FONT)
        add_text(slide, Emu(left_emu + int(Inches(0.20))),
                 top + Inches(1.30), Emu(int(col_w) - int(Inches(0.40))),
                 Inches(0.30),
                 f"{cur_curr} · {scenario_upside:+.1f}%", size=12, bold=True,
                 color=c, font=BODY_FONT)
        # thesis
        add_text(slide, Emu(left_emu + int(Inches(0.20))),
                 top + Inches(1.75), Emu(int(col_w) - int(Inches(0.40))),
                 Inches(1.10),
                 sd.get('thesis', '')[:200], size=10, color=INK,
                 italic=True, font=BODY_FONT)
        # assumptions
        add_text(slide, Emu(left_emu + int(Inches(0.20))),
                 top + Inches(2.95), Emu(int(col_w) - int(Inches(0.40))),
                 Inches(0.25),
                 "핵심 가정", size=10, bold=True, color=DEEP_FOREST,
                 font=HEADER_FONT)
        assumptions = sd.get('key_assumptions', [])[:5]
        for j, a in enumerate(assumptions):
            atop = top + Inches(3.20) + Inches(0.25) * j
            add_circle(slide, Emu(left_emu + int(Inches(0.20))),
                       atop + Inches(0.08), Inches(0.10), c)
            add_text(slide, Emu(left_emu + int(Inches(0.40))), atop,
                     Emu(int(col_w) - int(Inches(0.60))), Inches(0.25),
                     a[:90], size=9, color=INK, font=BODY_FONT)


def catalysts_slide(prs, deep_research, page, total):
    """Slide 8 — Catalyst timeline."""
    catalysts = deep_research.get('catalysts', [])[:5]
    slide = base_slide(prs, "07 · Catalysts",
                       f"향후 12-24개월 주요 catalyst {len(catalysts)}건",
                       "각 이벤트별 예상 주가 영향 + 설명 — + 상승 / - 하락 / ○ 중립",
                       page, total)
    type_emoji = {"earnings": "📊", "policy": "🏛", "product": "🚀",
                  "m&a": "🤝", "regulatory": "⚖", "macro": "🌍"}
    for i, c in enumerate(catalysts):
        top = BODY_TOP + Inches(0.10) + Inches(0.90) * i
        impact = c.get('expected_impact', '○')
        ic = BULL if impact == '+' else (BEAR if impact == '-' else NEU)
        add_rect(slide, BODY_LEFT, top, BODY_W, Inches(0.80), WHITE,
                 line=LIGHT_GRAY)
        add_rect(slide, BODY_LEFT, top, Inches(0.10), Inches(0.80), ic)
        # date badge
        add_rect(slide, BODY_LEFT + Inches(0.25), top + Inches(0.15),
                 Inches(1.50), Inches(0.50), DEEP_FOREST)
        add_text(slide, BODY_LEFT + Inches(0.25), top + Inches(0.15),
                 Inches(1.50), Inches(0.50),
                 c.get('date', ''), size=12, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=BODY_FONT)
        # event title
        add_text(slide, BODY_LEFT + Inches(1.95), top + Inches(0.10),
                 Inches(8.5), Inches(0.35),
                 c.get('event', '')[:90], size=12, bold=True, color=DEEP_FOREST,
                 font=HEADER_FONT)
        # description
        add_text(slide, BODY_LEFT + Inches(1.95), top + Inches(0.45),
                 Inches(8.5), Inches(0.30),
                 c.get('description', '')[:130], size=10, color=INK,
                 font=BODY_FONT)
        # impact badge right
        add_circle(slide, BODY_LEFT + BODY_W - Inches(0.85),
                   top + Inches(0.20), Inches(0.55), ic)
        add_text(slide, BODY_LEFT + BODY_W - Inches(0.85),
                 top + Inches(0.20), Inches(0.55), Inches(0.55),
                 impact, size=24, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=HEADER_FONT)


def risks_slide(prs, deep_research, page, total):
    """Slide 9 — Risk matrix."""
    risks = deep_research.get('risks', [])[:5]
    slide = base_slide(prs, "08 · Risk Matrix",
                       f"5대 종목별 risk (Probability × Impact)",
                       "각 risk의 발생확률·영향도·설명·완화방안 — 진한 색일수록 우선 모니터링",
                       page, total)
    color_map = {
        ("high", "high"): BEAR, ("high", "med"): RGBColor(0xF5, 0x7C, 0x00),
        ("med", "high"): RGBColor(0xF5, 0x7C, 0x00),
        ("med", "med"): GOLD, ("med", "low"): RGBColor(0xCD, 0xDC, 0x39),
        ("low", "high"): GOLD, ("low", "med"): RGBColor(0xCD, 0xDC, 0x39),
        ("low", "low"): BULL,
    }
    label = {"high": "高", "med": "中", "low": "低"}
    headers = ["리스크", "P", "I", "Severity", "설명", "완화 방안"]
    col_w = [Inches(3.0), Inches(0.45), Inches(0.45), Inches(0.80),
             Inches(4.43), Inches(3.20)]
    col_x = [BODY_LEFT]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    hdr_top = BODY_TOP + Inches(0.10)
    add_rect(slide, BODY_LEFT, hdr_top, BODY_W, Inches(0.40), DEEP_FOREST)
    for i, h in enumerate(headers):
        add_text(slide, col_x[i], hdr_top, col_w[i], Inches(0.40), h, size=10,
                 bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, font=BODY_FONT)
    for i, r in enumerate(risks):
        row_top = hdr_top + Inches(0.45) + Inches(0.80) * i
        if i % 2 == 0:
            add_rect(slide, BODY_LEFT, row_top, BODY_W, Inches(0.75), WHITE)
        p = r.get('probability', 'med')
        ipa = r.get('impact', 'med')
        sev = color_map.get((p, ipa), NEU)
        add_text(slide, col_x[0] + Inches(0.10), row_top, col_w[0],
                 Inches(0.75), r.get('name', '')[:60], size=10, bold=True,
                 color=DEEP_FOREST, anchor=MSO_ANCHOR.MIDDLE,
                 font=BODY_FONT)
        for j, v in enumerate([p, ipa]):
            add_text(slide, col_x[1+j], row_top, col_w[1+j], Inches(0.75),
                     label.get(v, '?'), size=14, bold=True, color=sev,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                     font=BODY_FONT)
        add_circle(slide, col_x[3] + Inches(0.27), row_top + Inches(0.20),
                   Inches(0.30), sev)
        add_text(slide, col_x[4] + Inches(0.05), row_top, col_w[4],
                 Inches(0.75), r.get('description', '')[:120], size=9,
                 color=INK, anchor=MSO_ANCHOR.MIDDLE, font=BODY_FONT)
        add_text(slide, col_x[5] + Inches(0.05), row_top, col_w[5],
                 Inches(0.75), r.get('mitigation', '')[:90], size=9,
                 color=MOSS, italic=True, anchor=MSO_ANCHOR.MIDDLE,
                 font=BODY_FONT)


def industry_slide(prs, deep_research, page, total):
    """Slide 10 — Industry context + competitors."""
    ind = deep_research.get('industry', {})
    competitors = ind.get('competitors', [])[:5]
    slide = base_slide(prs, "09 · Industry Context",
                       f"{ind.get('name', '-')}",
                       f"시장 규모 ${ind.get('market_size', {}).get('current_usd_bn', 0)}B · CAGR {ind.get('market_size', {}).get('cagr_pct', 0)}% · {ind.get('market_size', {}).get('horizon', '')}",
                       page, total)
    headers = ["회사", "Ticker", "점유율", "Moat (경쟁우위)"]
    col_w = [Inches(2.6), Inches(1.4), Inches(1.0), Inches(7.33)]
    col_x = [BODY_LEFT]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    hdr_top = BODY_TOP + Inches(0.10)
    add_rect(slide, BODY_LEFT, hdr_top, BODY_W, Inches(0.40), DEEP_FOREST)
    for i, h in enumerate(headers):
        add_text(slide, col_x[i], hdr_top, col_w[i], Inches(0.40), h, size=11,
                 bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, font=BODY_FONT)
    for i, c in enumerate(competitors):
        row_top = hdr_top + Inches(0.45) + Inches(0.80) * i
        if i % 2 == 0:
            add_rect(slide, BODY_LEFT, row_top, BODY_W, Inches(0.75), WHITE)
        add_text(slide, col_x[0] + Inches(0.10), row_top, col_w[0],
                 Inches(0.75), c.get('name', '')[:30], size=11, bold=True,
                 color=DEEP_FOREST, anchor=MSO_ANCHOR.MIDDLE,
                 font=BODY_FONT)
        add_text(slide, col_x[1], row_top, col_w[1], Inches(0.75),
                 c.get('ticker', '-'), size=10, color=GRAY,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=BODY_FONT)
        # share bar
        share = c.get('share_pct', 0)
        bar_full = int(col_w[2]) - int(Inches(0.20))
        add_rect(slide, col_x[2] + Inches(0.10), row_top + Inches(0.30),
                 Emu(bar_full), Inches(0.15), LIGHT_GRAY)
        add_rect(slide, col_x[2] + Inches(0.10), row_top + Inches(0.30),
                 Emu(int(bar_full * share / 40)), Inches(0.15), GOLD)
        add_text(slide, col_x[2], row_top + Inches(0.45), col_w[2],
                 Inches(0.30), f"{share:.1f}%", size=9, color=GRAY,
                 align=PP_ALIGN.CENTER, font=BODY_FONT)
        add_text(slide, col_x[3] + Inches(0.10), row_top, col_w[3],
                 Inches(0.75), c.get('moat', '')[:100], size=10, color=INK,
                 anchor=MSO_ANCHOR.MIDDLE, italic=True, font=BODY_FONT)


def decision_slide(prs, ticker, decision, persona_agg, page, total):
    """Slide 11 — Final decision sheet (dark dramatic finish)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = DEEP_FOREST
    bg.line.fill.background()

    add_rect(slide, Inches(0.7), Inches(0.7), Inches(2.0), Inches(0.04), GOLD)
    add_text(slide, Inches(0.7), Inches(0.85), Inches(11.9), Inches(0.4),
             "10 · DECISION SHEET", size=11, bold=True, color=GOLD,
             font=BODY_FONT)

    # Action big
    action = decision.get('action', '-')
    add_text(slide, Inches(0.7), Inches(1.7), Inches(11.9), Inches(1.2),
             action, size=72, bold=True,
             color=BULL if 'BUY' in action else GOLD, font=HEADER_FONT)

    # Sub
    add_text(slide, Inches(0.7), Inches(2.95), Inches(11.9), Inches(0.5),
             decision.get('rationale', '')[:140], size=14, color=SAGE,
             italic=True, font=BODY_FONT)

    # 3 KPI strip
    kpis = [
        ("Position size", f"{decision.get('size_pct', 0):.1f}%",
         f"~${decision.get('size_usd', 0):,} / $100K"),
        ("Score", f"{persona_agg.get('weighted_average_score_1to5', 0):.2f}",
         f"평균 신뢰도 {int(persona_agg.get('average_confidence', 0)*100)}%"),
        ("Horizon", decision.get('horizon', '-')[:20], ""),
    ]
    kpi_top = Inches(4.2)
    kpi_w = Inches(3.95)
    for i, (lbl, val, sub) in enumerate(kpis):
        kx = Inches(0.7) + (kpi_w + Inches(0.10)) * i
        add_rect(slide, kx, kpi_top, kpi_w, Inches(1.6), MOSS)
        add_text(slide, kx + Inches(0.20), kpi_top + Inches(0.15),
                 kpi_w - Inches(0.40), Inches(0.25), lbl, size=10,
                 color=SAGE, font=BODY_FONT)
        add_text(slide, kx + Inches(0.20), kpi_top + Inches(0.40),
                 kpi_w - Inches(0.40), Inches(0.75), val, size=28, bold=True,
                 color=GOLD, font=HEADER_FONT)
        add_text(slide, kx + Inches(0.20), kpi_top + Inches(1.20),
                 kpi_w - Inches(0.40), Inches(0.30), sub, size=10,
                 color=SAGE, font=BODY_FONT)

    # Stop loss + add trigger
    add_text(slide, Inches(0.7), Inches(6.10), Inches(11.9), Inches(0.30),
             "  Risk Management", size=11, bold=True, color=GOLD,
             font=HEADER_FONT)
    add_text(slide, Inches(0.7), Inches(6.40), Inches(5.9), Inches(0.40),
             f"Stop loss: {decision.get('stop_loss', '-')[:80]}",
             size=11, color=SAGE, font=BODY_FONT)
    add_text(slide, Inches(6.7), Inches(6.40), Inches(5.9), Inches(0.40),
             f"Add trigger: {decision.get('add_trigger', '-')[:80]}",
             size=11, color=SAGE, font=BODY_FONT)

    # footer
    add_text(slide, Inches(0.7), Inches(7.05), Inches(11.9), Inches(0.30),
             f"{ticker} · 13명 페르소나 + Deep Research · 분석 2026-05-07",
             size=9, color=SAGE, font=BODY_FONT)
    add_text(slide, Inches(0.7), Inches(7.05), Inches(11.9), Inches(0.30),
             f"{page:02d} / {total:02d}", size=9, color=SAGE,
             align=PP_ALIGN.RIGHT, font=BODY_FONT)


def build_pptx_for_ticker(pipeline_dir: Path, ticker: str, output: Path):
    meta = _read_json(pipeline_dir / "meta.json")
    stocks_raw = _read_json(pipeline_dir / "stocks.json")
    persona_agg = _read_json(pipeline_dir / "persona_panel" / ticker / "aggregate.json")
    deep_research = _read_json(pipeline_dir / "deep_research" / f"{ticker}.json")
    decisions_data = _read_json(pipeline_dir / "decisions.json").get("decisions", [])

    # Find ticker entries
    market = {}
    company = ticker
    for s in stocks_raw:
        if s.get('ticker') == ticker:
            market = s.get('market_data', {})
            company = s.get('name', ticker)
            break
    decision = next((d for d in decisions_data if d.get('ticker') == ticker), {})
    score = persona_agg.get('weighted_average_score_1to5', 0)
    action = decision.get('action', '-')

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    TOTAL = 11
    cover_slide(prs, ticker, company, market, score, action)
    snapshot_slide(prs, ticker, company, market, score, 2, TOTAL)
    thesis_alignment_slide(prs, deep_research, 3, TOTAL)
    persona_verdict_slide(prs, persona_agg, 4, TOTAL)
    style_split_slide(prs, persona_agg, 5, TOTAL)
    concerns_opportunities_slide(prs, persona_agg, 6, TOTAL)
    scenarios_slide(prs, deep_research, market, 7, TOTAL)
    catalysts_slide(prs, deep_research, 8, TOTAL)
    risks_slide(prs, deep_research, 9, TOTAL)
    industry_slide(prs, deep_research, 10, TOTAL)
    decision_slide(prs, ticker, decision, persona_agg, 11, TOTAL)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    print(f"[per-stock] saved {output} ({len(prs.slides)} slides)")


def main():
    parser = argparse.ArgumentParser(description="Per-stock PPT builder")
    parser.add_argument("--pipeline-dir", required=True)
    parser.add_argument("--ticker", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    build_pptx_for_ticker(Path(args.pipeline_dir), args.ticker, Path(args.output))


if __name__ == "__main__":
    main()
